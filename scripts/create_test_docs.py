import os
import io
import base64
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Sample 10x10 PNG Base64
SAMPLE_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAAoAAAAKCAYAAACNMs+9AAAAFUlEQVR42mNk+M9Qz0AEYBhxVYACAB75C/365Dq4AAAAAElFTkSuQmCC"

def get_image_stream():
    return io.BytesIO(base64.b64decode(SAMPLE_PNG_B64))

def add_hyperlink(paragraph, url, text, color="0066CC", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    xml_str = (
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:color w:val="{color}"/>'
        f'<w:u w:val="{"single" if underline else "none"}"/>'
        f'</w:rPr>'
        f'<w:t>{text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    hyperlink = parse_xml(xml_str)
    paragraph._p.append(hyperlink)

def create_01_basic_structure(output_dir):
    doc = docx.Document()
    doc.add_heading('Blog Title Example: Basic Structure', level=0)
    
    doc.add_heading('Main Heading Level 1', level=1)
    doc.add_paragraph('This is a standard body paragraph under Heading 1 explaining the primary topic.')
    
    doc.add_heading('Sub-Heading Level 2', level=2)
    doc.add_paragraph('This paragraph resides under Heading 2 and provides secondary details.')
    
    doc.add_heading('Detailed Section Level 3', level=3)
    doc.add_paragraph('Final paragraph describing nested document structures and subsection depth.')
    
    filepath = os.path.join(output_dir, '01-basic-structure.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_02_formatting(output_dir):
    doc = docx.Document()
    doc.add_heading('Typography & Formatting Test', level=1)
    
    p1 = doc.add_paragraph('This paragraph contains regular text, ')
    p1.add_run('bold text').bold = True
    p1.add_run(', ')
    p1.add_run('italic text').italic = True
    p1.add_run(', and ')
    r_bi = p1.add_run('bold italic text')
    r_bi.bold = True
    r_bi.italic = True
    p1.add_run(' inline within a single sentence.')
    
    p2 = doc.add_paragraph('Another paragraph testing inline styling variations like ')
    p2.add_run('strong emphasis').bold = True
    p2.add_run(' and ')
    p2.add_run('soft emphasis').italic = True
    p2.add_run('.')
    
    filepath = os.path.join(output_dir, '02-formatting.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_03_lists(output_dir):
    doc = docx.Document()
    doc.add_heading('Lists & Bullet Items Test', level=1)
    
    doc.add_paragraph('Unordered Bullet List:')
    doc.add_paragraph('First unordered bullet item', style='List Bullet')
    doc.add_paragraph('Second unordered bullet item', style='List Bullet')
    doc.add_paragraph('Third unordered bullet item', style='List Bullet')
    
    doc.add_paragraph('Ordered Numbered List:')
    doc.add_paragraph('First ordered numbered step', style='List Number')
    doc.add_paragraph('Second ordered numbered step', style='List Number')
    doc.add_paragraph('Third ordered numbered step', style='List Number')
    
    filepath = os.path.join(output_dir, '03-lists.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_04_links(output_dir):
    doc = docx.Document()
    doc.add_heading('Hyperlinks Test', level=1)
    
    p1 = doc.add_paragraph('For more details, visit the official site at ')
    add_hyperlink(p1, 'https://www.postgresql.org/', 'PostgreSQL Official Website')
    p1.add_run(' for release notes and documentation.')
    
    p2 = doc.add_paragraph('Check out ')
    add_hyperlink(p2, 'https://graphql.org/', 'GraphQL Documentation')
    p2.add_run(' for API specifications.')
    
    filepath = os.path.join(output_dir, '04-links.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_05_images(output_dir):
    doc = docx.Document()
    doc.add_heading('Embedded Images Test', level=1)
    
    doc.add_paragraph('Below is the first embedded demonstration image:')
    doc.add_picture(get_image_stream(), width=Inches(2.0))
    p1 = doc.add_paragraph('Figure 1: Primary Test Diagram Asset')
    p1.runs[0].italic = True
    
    doc.add_paragraph('Below is the second embedded demonstration image:')
    doc.add_picture(get_image_stream(), width=Inches(2.5))
    p2 = doc.add_paragraph('Figure 2: Secondary Test Flowchart Asset')
    p2.runs[0].italic = True
    
    filepath = os.path.join(output_dir, '05-images.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_06_blockquotes_callouts(output_dir):
    doc = docx.Document()
    doc.add_heading('Quotes and Callouts Test', level=1)
    
    doc.add_paragraph('Standard body context prior to blockquote.')
    
    doc.add_paragraph('“The secret of getting ahead is getting started. The secret of getting started is breaking your complex overwhelming tasks into small manageable tasks.”', style='Quote')
    
    p_callout = doc.add_paragraph()
    r_h = p_callout.add_run('IMPORTANT CALLOUT: ')
    r_h.bold = True
    p_callout.add_run('This is a highlighted callout box section containing essential requirements and architecture instructions for the engineering team.')
    
    filepath = os.path.join(output_dir, '06-blockquotes-callouts.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_07_tables(output_dir):
    doc = docx.Document()
    doc.add_heading('Structured Data Tables Test', level=1)
    
    doc.add_paragraph('Below is a comparison table of application stack layers:')
    
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Feature / Technology'
    hdr_cells[1].text = 'Status'
    hdr_cells[2].text = 'Description'
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                
    row1 = table.rows[1].cells
    row1[0].text = 'Next.js'
    row1[1].text = 'Active'
    row1[2].text = 'Frontend App Router Framework'
    
    row2 = table.rows[2].cells
    row2[0].text = 'GraphQL / Apollo'
    row2[1].text = 'Active'
    row2[2].text = 'Strongly-typed API Layer'
    
    row3 = table.rows[3].cells
    row3[0].text = 'PostgreSQL / Prisma'
    row3[1].text = 'Active'
    row3[2].text = 'Relational Persistence Layer'
    
    filepath = os.path.join(output_dir, '07-tables.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_08_combined_blog(output_dir):
    doc = docx.Document()
    
    # Title
    doc.add_heading('Comprehensive Guide to Modern Web Architecture', level=0)
    
    # Author / Date Metadata
    meta = doc.add_paragraph('By Technical Lead | Published September 2026')
    meta.runs[0].italic = True
    
    # Featured Image
    doc.add_picture(get_image_stream(), width=Inches(3.5))
    cap = doc.add_paragraph('Featured Banner Architecture Diagram')
    cap.runs[0].italic = True
    
    # Section 1
    doc.add_heading('1. Introduction to Monorepos', level=1)
    p1 = doc.add_paragraph('Building modern enterprise web applications requires ')
    p1.add_run('strong architectural decoupling').bold = True
    p1.add_run(' between frontend client interfaces and backend GraphQL services.')
    
    # Section 2
    doc.add_heading('2. Key Architecture Benefits', level=2)
    doc.add_paragraph('Single repository for frontend and backend codebases.', style='List Bullet')
    doc.add_paragraph('Unified TypeScript types and end-to-end type safety.', style='List Bullet')
    doc.add_paragraph('Streamlined local development and testing workflows.', style='List Bullet')
    
    # Section 3
    doc.add_heading('3. Implementation Steps', level=2)
    doc.add_paragraph('Initialize root npm workspace package.', style='List Number')
    doc.add_paragraph('Configure Apollo GraphQL server and Prisma ORM.', style='List Number')
    doc.add_paragraph('Connect Next.js Apollo Client to Apollo Server endpoint.', style='List Number')
    
    # Blockquote
    doc.add_paragraph('“Good software architecture makes system evolution predictable, maintainable, and resilient to change.”', style='Quote')
    
    # Link
    p_link = doc.add_paragraph('For architectural standards, refer to ')
    add_hyperlink(p_link, 'https://graphql.org/', 'Official GraphQL Standards')
    p_link.add_run('.')
    
    # Table
    doc.add_heading('4. Layer Summary Table', level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Layer'
    hdr[1].text = 'Technology'
    hdr[2].text = 'Role'
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                
    r1 = table.rows[1].cells
    r1[0].text = 'Frontend'
    r1[1].text = 'Next.js 15'
    r1[2].text = 'UI Component Rendering'
    
    r2 = table.rows[2].cells
    r2[0].text = 'Backend'
    r2[1].text = 'Apollo Server'
    r2[2].text = 'GraphQL Resolver API'
    
    # Conclusion
    doc.add_heading('5. Conclusion', level=2)
    doc.add_paragraph('Following these principles ensures robust scalability and maintainability across team boundaries.')
    
    filepath = os.path.join(output_dir, '08-combined-blog.docx')
    doc.save(filepath)
    print(f"Created: {filepath}")

if __name__ == '__main__':
    out_dir = os.path.join(os.getcwd(), 'sample-documents', 'parser-tests')
    os.makedirs(out_dir, exist_ok=True)
    
    create_01_basic_structure(out_dir)
    create_02_formatting(out_dir)
    create_03_lists(out_dir)
    create_04_links(out_dir)
    create_05_images(out_dir)
    create_06_blockquotes_callouts(out_dir)
    create_07_tables(out_dir)
    create_08_combined_blog(out_dir)
    print("All 8 test fixtures created successfully!")
